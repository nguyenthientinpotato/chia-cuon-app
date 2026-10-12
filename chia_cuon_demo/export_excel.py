"""Xuat ket qua ra Excel theo format block cu tren sheet FastCal (moi lenh 1 block)."""
from __future__ import annotations

import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import slitter as s

HEAD = PatternFill("solid", fgColor="DDE7F3")
WARN = PatternFill("solid", fgColor="F8D7D6")
OKF = PatternFill("solid", fgColor="DCEFE3")
THIN = Side(style="thin", color="B7B7B7")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)


def _row(ws, r, values, bold=False, fill=None, fmt=None):
    for c, v in enumerate(values, start=1):
        cell = ws.cell(r, c, v)
        cell.border = BOX
        if bold:
            cell.font = BOLD
        if fill:
            cell.fill = fill
        if fmt and isinstance(v, (int, float)) and not isinstance(v, bool):
            cell.number_format = fmt.get(c, "General")


def to_excel_bytes(sol: s.Solution, cfg: dict | None = None) -> bytes:
    p = sol.problem
    wb = Workbook()

    # ---- Sheet 1: lenh chia (block giong FastCal) ----
    ws = wb.active
    ws.title = "Lenh chia"
    t = s.totals(sol)
    multi = len(p.masters) > 1
    ws["A1"], ws["B1"] = "Kho NL (mm)", ", ".join(f"{m.width_mm:g}" for m in p.masters) if multi else p.master_width_mm
    ws["C1"], ws["D1"] = "Dinh luong (g/m2)", ", ".join(f"{m.gsm:g}" for m in p.masters) if multi else p.gsm
    ws["E1"], ws["F1"] = "Tong NVL (kg)", t["NVL (kg)"]
    ws["G1"], ws["H1"] = "Phe (kg)", t["Phe (kg)"]
    ws["I1"], ws["J1"] = "% phe", t["% phe"] / 100
    ws["J1"].number_format = "0.00%"
    for c in "ACEGI":
        ws[f"{c}1"].font = BOLD

    hdr = ["Lenh", "Kho NL", "SL NL can (kg)", "So m/bo", "So bo", "Kho", "SD", "Loai",
           "Cho don", "SC tren may", "SC thanh pham", "Noi cuon", "So kg"]
    fmt = {3: "#,##0.0", 13: "#,##0.0"}
    r = 3
    plan = s.plan_rows(sol)
    for n, (pt, sets) in enumerate(sol.runs, start=1):
        rows = [x for x in plan if x["Lenh"] == n]
        master_w = rows[0]["Kho NL"]
        nl_tag = f" [{rows[0]['Cuon NL']} - {master_w:g} mm]" if multi else ""
        ws.cell(r, 1, f"Lenh {n}{nl_tag}: {rows[0]['Cong thuc']}").font = Font(bold=True, size=12)
        r += 1
        _row(ws, r, hdr, bold=True, fill=HEAD)
        r += 1
        for i, x in enumerate(rows):
            first = i == 0
            _row(ws, r, [n if first else None, master_w if first else None,
                         x["SL NL can (kg)"] if first else None, pt.length_m if first else None,
                         sets if first else None, x["Kho cuon con"], x["SD"], x["Loai"], x["Don"],
                         x["SC tren may"], x["SC thanh pham"], x["Noi cuon"], x["So kg"]], fmt=fmt)
            r += 1
        trim_mm, trim_kg = rows[0]["Bien phe (mm)"], rows[0]["Phe (kg)"]
        _row(ws, r, ["Phe lieu", None, trim_kg, None, None, trim_mm, None, None,
                     f"{trim_mm / master_w:.2%}", None, None, None, None], bold=True, fmt=fmt)
        r += 2
    for c, w in enumerate([12, 10, 15, 10, 8, 9, 6, 8, 22, 12, 14, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # ---- Sheet 2: dap ung theo don ----
    ws2 = wb.create_sheet("Dap ung")
    summ = s.summary_rows(sol)
    cols = list(summ[0].keys())
    _row(ws2, 1, cols, bold=True, fill=HEAD)
    for i, x in enumerate(summ, start=2):
        st = x["Trang thai"]
        fill = WARN if st in ("VUOT DUNG SAI", "THIEU", "DU (VUA DU)") else (OKF if st == "OK" else None)
        _row(ws2, i, [x[c] for c in cols], fill=fill)
    for c in range(1, len(cols) + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 16

    # ---- Sheet NVL: muc su dung tung cuon NVL cap ----
    ws_nl = wb.create_sheet("NVL")
    nl = s.master_rows(sol)
    _row(ws_nl, 1, list(nl[0].keys()), bold=True, fill=HEAD)
    for i, x in enumerate(nl, start=2):
        _row(ws_nl, i, list(x.values()))
    for c in range(1, len(nl[0]) + 1):
        ws_nl.column_dimensions[get_column_letter(c)].width = 14

    # ---- Sheet 3: tong hop + input (de luu vet) ----
    ws3 = wb.create_sheet("Tong hop")
    _row(ws3, 1, ["Chi so", "Gia tri"], bold=True, fill=HEAD)
    rr = 2
    for k, v in t.items():
        _row(ws3, rr, [k, v])
        rr += 1
    _row(ws3, rr, ["Trang thai solver", sol.status])
    _row(ws3, rr + 1, ["Phuong phap", sol.method])
    if cfg:
        ws3.cell(rr + 3, 1, "Input (JSON)").font = BOLD
        c = ws3.cell(rr + 4, 1, json.dumps(cfg, ensure_ascii=False, indent=1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
